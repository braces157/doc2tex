"""Create a deterministic synthetic Word bookmark/field fixture for R07."""
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1] / 'src/test/resources/corpus'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
doc = Document()
doc.core_properties.title = 'Synthetic caption and explicit reference fixture'
doc.core_properties.author = 'doc2TeX synthetic fixtures'
doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 16, tzinfo=timezone.utc)
doc.add_heading('Caption and bookmark reference checks', 0)
doc.add_paragraph('PRESERVE_REFERENCE_START. Only explicit supported Word reference fields should become numbered LaTeX references.')
with ZipFile(ROOT / 'research-complex.docx') as source:
    image = source.read('word/media/image1.png')

next_bookmark = 1


def simple_field(paragraph, instruction, cached):
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = cached
    run.append(text); field.append(run); paragraph._p.append(field)


def complex_field(paragraph, instruction, cached):
    for field_type in ['begin', 'instruction', 'separate', 'result', 'end']:
        run = OxmlElement('w:r')
        if field_type == 'instruction':
            text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction; run.append(text)
        elif field_type == 'result':
            text = OxmlElement('w:t'); text.text = cached; run.append(text)
        else:
            char = OxmlElement('w:fldChar'); char.set(qn('w:fldCharType'), field_type); run.append(char)
        paragraph._p.append(run)


def caption(kind, number, bookmark, text):
    global next_bookmark
    paragraph = doc.add_paragraph(style='Caption'); paragraph.add_run(kind + ' ')
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(next_bookmark)); start.set(qn('w:name'), bookmark); paragraph._p.append(start)
    simple_field(paragraph, f' SEQ {kind} \\* ARABIC ', str(number))
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(next_bookmark)); paragraph._p.append(end)
    next_bookmark += 1
    paragraph.add_run(' ' + text)


for number, bookmark in [(1, '_RefFigureAlpha'), (2, '_RefFigureBeta')]:
    doc.add_picture(BytesIO(image), width=Inches(2.5))
    caption('Figure', number, bookmark, 'Repeated sensor observations')

caption('Table', 1, '_RefTableAlpha', 'Measured reference outcomes')
table = doc.add_table(rows=2, cols=2); table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['Condition', 'Result']): cell.text = text
for cell, text in zip(table.rows[1].cells, ['Reference measurement', 'PRESERVE_REFERENCE_TABLE']): cell.text = text

p = doc.add_paragraph('REFERENCE_FIGURE_ALPHA='); simple_field(p, ' REF _RefFigureAlpha \\h ', '1'); p.add_run('.')
p = doc.add_paragraph('REFERENCE_FIGURE_BETA='); complex_field(p, ' REF _RefFigureBeta \\h ', '2'); p.add_run('.')
p = doc.add_paragraph('REFERENCE_TABLE_ALPHA='); simple_field(p, ' REF _RefTableAlpha \\h ', '1'); p.add_run('.')
doc.add_paragraph('PRESERVE_REFERENCE_PROSE: see Figure 1 and Table 2. These ordinary words have no reference relationship.')
p = doc.add_paragraph('UNSUPPORTED_PAGE_REFERENCE='); simple_field(p, ' PAGEREF _RefFigureAlpha \\h ', '1'); p.add_run('.')
p = doc.add_paragraph('UNRESOLVED_REFERENCE='); simple_field(p, ' REF MissingBookmark \\h ', 'UNRESOLVED_REFERENCE'); p.add_run('.')
doc.add_paragraph('A detached caption follows without an adjacent figure or table.')
caption('Figure', 3, '_RefOrphan', 'Detached caption requires manual ownership review')
p = doc.add_paragraph('AMBIGUOUS_OWNER_REFERENCE='); simple_field(p, ' REF _RefOrphan \\h ', '3'); p.add_run('.')
doc.add_paragraph('PRESERVE_REFERENCE_FINISH')

buffer = BytesIO(); doc.save(buffer)
with ZipFile(buffer) as generated, ZipFile(ROOT / 'caption-references.docx', 'w', ZIP_DEFLATED) as output:
    for name in sorted(generated.namelist()):
        info = ZipInfo(name, (2026, 10, 16, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
        output.writestr(info, generated.read(name))

manifest_path = ROOT / 'manifest.json'
manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
manifest = [case for case in manifest if case['file'] != 'caption-references.docx']
manifest.append({'file': 'caption-references.docx',
                 'mustContain': ['PRESERVE_REFERENCE_START', 'PRESERVE_REFERENCE_TABLE', 'PRESERVE_REFERENCE_PROSE', 'PRESERVE_REFERENCE_FINISH'],
                 'features': ['numbered captions', 'repeated caption text', 'Word bookmarks', 'simple and complex REF fields', 'PAGEREF', 'unresolved and detached references'],
                 'previewExpected': 'caption and reference review required',
                 'pdfMustContain': ['PRESERVE_REFERENCE_START', 'PRESERVE_REFERENCE_TABLE', 'PRESERVE_REFERENCE_PROSE', 'PRESERVE_REFERENCE_FINISH'],
                 'pdfMustNotContain': [], 'pdfMinPages': 1})
manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8', newline='\n')
print('Created caption-references.docx and its corpus manifest entry.')
